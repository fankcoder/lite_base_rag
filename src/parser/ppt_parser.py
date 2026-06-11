"""
PPT 文档解析器
"""
import sys
import os
from pathlib import Path
from typing import List

BASE_DIR = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(BASE_DIR))

from .base_parser import BaseParser, DocumentChunk


class PPTParser(BaseParser):
    """PPTX 文档解析器（基于 python-pptx）"""

    def __init__(self, chunk_size=500, chunk_overlap=100,
                 extract_notes=True, extract_images=False,
                 generate_thumbnails=False):
        super().__init__(chunk_size, chunk_overlap)
        self.extract_notes = extract_notes
        self.extract_images = extract_images
        self.generate_thumbnails = generate_thumbnails

    def parse(self, file_path: str) -> List[DocumentChunk]:
        """解析 PPTX 文件"""
        from pptx import Presentation

        file_name = os.path.basename(file_path)
        file_hash = self.get_file_hash(file_path)

        prs = Presentation(file_path)
        chunks = []
        chunk_idx = 0

        for slide_num, slide in enumerate(prs.slides, start=1):
            # 提取页面文本
            slide_texts = []
            slide_title = ""

            for shape in slide.shapes:
                # 标题
                if shape.has_text_frame and shape == slide.shapes.title:
                    slide_title = shape.text_frame.text.strip()

                # 普通文本框
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)

                # 表格
                if shape.has_table:
                    table_md = self._table_to_markdown(shape.table)
                    if table_md:
                        # 表格单独作为 chunk
                        enriched = f"【文档】{file_name}\n【页码】第 {slide_num} 页"
                        if slide_title:
                            enriched += f"\n【页面】{slide_title}"
                        enriched += f"\n【表格】\n{table_md}"

                        chunk = DocumentChunk(
                            file_path=file_path,
                            file_name=file_name,
                            file_type='pptx',
                            file_hash=file_hash,
                            content=enriched,
                            content_type='table',
                            slide_num=slide_num,
                            chapter_path=slide_title,
                            chunk_index=chunk_idx,
                        )
                        chunks.append(chunk)
                        chunk_idx += 1

            # 提取演讲者备注
            notes_text = ""
            if self.extract_notes and slide.has_notes_slide:
                notes_slide = slide.notes_slide
                notes_text = notes_slide.notes_text_frame.text.strip()

            # 页面文本合并
            full_text = '\n'.join(slide_texts)

            if full_text.strip() or notes_text:
                # 组合页面内容
                content_parts = []
                if full_text.strip():
                    content_parts.append(full_text.strip())
                if notes_text:
                    content_parts.append(f"【备注】{notes_text}")

                combined_text = '\n\n'.join(content_parts)

                # 注入上下文
                enriched = f"【文档】{file_name}\n【页码】第 {slide_num} 页"
                if slide_title:
                    enriched += f"\n【页面标题】{slide_title}"
                enriched += f"\n【内容】\n{combined_text}"

                # 文本切片（PPT 每页内容一般不多，整页作为一个 chunk 也可以）
                if len(combined_text) > self.chunk_size * 2:
                    # 内容多的话再切
                    text_chunks = self.split_text(combined_text)
                    for tc in text_chunks:
                        chunk_enriched = f"【文档】{file_name}\n【页码】第 {slide_num} 页"
                        if slide_title:
                            chunk_enriched += f"\n【页面标题】{slide_title}"
                        chunk_enriched += f"\n【内容】\n{tc}"

                        chunk = DocumentChunk(
                            file_path=file_path,
                            file_name=file_name,
                            file_type='pptx',
                            file_hash=file_hash,
                            content=chunk_enriched,
                            content_type='text',
                            slide_num=slide_num,
                            chapter_path=slide_title,
                            chunk_index=chunk_idx,
                        )
                        chunks.append(chunk)
                        chunk_idx += 1
                else:
                    chunk = DocumentChunk(
                        file_path=file_path,
                        file_name=file_name,
                        file_type='pptx',
                        file_hash=file_hash,
                        content=enriched,
                        content_type='text',
                        slide_num=slide_num,
                        chapter_path=slide_title,
                        chunk_index=chunk_idx,
                    )
                    chunks.append(chunk)
                    chunk_idx += 1

        return chunks

    def _table_to_markdown(self, table) -> str:
        """将 PPT 表格转为 markdown 格式"""
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cells.append(cell.text.strip())
            rows.append(cells)

        if len(rows) < 2:
            return ""

        max_cols = max(len(row) for row in rows)
        rows = [list(row) + [''] * (max_cols - len(row)) for row in rows]

        lines = []
        lines.append('| ' + ' | '.join(rows[0]) + ' |')
        lines.append('| ' + ' | '.join(['---'] * max_cols) + ' |')
        for row in rows[1:]:
            lines.append('| ' + ' | '.join(row) + ' |')

        return '\n'.join(lines)

    def _extract_images(self, slide, output_dir: str, slide_num: int) -> List[str]:
        """提取 PPT 中的图片（可选）"""
        image_paths = []
        for i, shape in enumerate(slide.shapes):
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    image = shape.image
                    image_bytes = image.blob
                    ext = image.content_type.split('/')[-1]
                    image_path = os.path.join(output_dir, f"slide_{slide_num}_img_{i}.{ext}")
                    with open(image_path, 'wb') as f:
                        f.write(image_bytes)
                    image_paths.append(image_path)
                except Exception:
                    pass
        return image_paths
